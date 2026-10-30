"""Spacious 1223-inspired investor report: five requested sections."""
import io
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

import streamlit as st

st.set_page_config(page_title="Moo Housing", layout="wide")

st.title("Moo Housing Investor Report Generator")
st.success("Application loaded successfully!")

BLUE='203D57'; MUTED='66798A'
def num(v):
    try: return float(v or 0)
    except (TypeError,ValueError): return 0.0
def money(v): return f'${num(v):,.2f}'
def value(row,*keys):
    for key in keys:
        v=row.get(key)
        if v is not None and str(v).strip()!='': return v
    return None
def percent(v):
    try:
        f=float(v)
        return f'{f*100:.1f}%' if f<=1 else f'{f:.1f}%'
    except (ValueError,TypeError): return 'Not provided'
def content(row,p,cn):
    occupancy=value(row,'Occupancy')
    occupied=value(row,'Occupied Rooms'); total=value(row,'Total Rooms')
    vacant=(int(num(total))-int(num(occupied))) if occupied is not None and total is not None else None
    occ=percent(occupancy) if occupancy is not None else ('未提供' if cn else 'Not provided')
    rent=value(row,'Rent Collected','Rent Collection','Collected Rent')
    ar=value(row,'Outstanding AR','Accounts Receivable','AR')
    cash_verified=value(row,'Verified Ending Bank Balance')
    summary=(f"本期出租率为 {occ}，净运营收入（NOI）为 {money(row.get('NOI'))}。" if cn else f"The reporting period recorded {occ} occupancy and {money(row.get('NOI'))} in net operating income (NOI).")
    if num(row.get('Reserve Shortfall'))>0:
        summary+=(' 期末存在储备缺口，建议优先补足后再考虑分红。' if cn else ' A reserve shortfall remains; replenishment should take priority over distributions.')
    else: summary+=(' 期末无储备缺口，分红仍需核对银行及应付款项。' if cn else ' No reserve shortfall is indicated; distributions remain subject to bank and payable verification.')
    notes=str(row.get('Notes') or '').strip()
    if notes: summary+='\n'+notes
    sections=[
        ('1. Summary' if not cn else '1. Summary｜经营摘要',summary,[]),
        ('2. Occupancy & Rent Collection' if not cn else '2. Occupancy & Rent Collection｜出租与收租',
         ('出租率按 AppFolio Rent Roll 计算；整租主行与子房间不可重复计数。' if cn else 'Occupancy is calculated from the AppFolio rent roll; whole-unit master leases and underlying rooms must not be double-counted.'),
         [('Occupancy Rate' if not cn else '出租率',occ),('Occupied Rooms' if not cn else '已出租房间',str(int(num(occupied))) if occupied is not None else 'Not provided'),('Vacant Rooms' if not cn else '空置房间',str(vacant) if vacant is not None else 'Not provided'),('Total Rooms' if not cn else '总房间',str(int(num(total))) if total is not None else 'Not provided'),('Rent Collected' if not cn else '已收租金',money(rent) if rent is not None else ('未提供' if cn else 'Not provided')),('Outstanding AR' if not cn else '未收应收租金',money(ar) if ar is not None else ('未提供' if cn else 'Not provided'))]),
        ('3. Financial Summary' if not cn else '3. Financial Summary｜财务摘要',
         ('NOI 为会计经营指标，不直接等于银行现金余额。' if cn else 'NOI is an operating accounting measure and does not necessarily equal bank cash.'),
         [('Total Income' if not cn else '总收入',money(row.get('Income'))),('Operating Expenses' if not cn else '运营费用',money(row.get('Expenses'))),('Net Operating Income (NOI)' if not cn else '净运营收入（NOI）',money(row.get('NOI'))),('Ending Cash Balance' if not cn else '期末现金余额',money(row.get('Ending Cash'))),('Required Reserve' if not cn else '累计应保留储备',money(row.get('Required Reserve')))] + ([('Reserve Shortfall' if not cn else '储备缺口',money(row.get('Reserve Shortfall')))] if num(row.get('Reserve Shortfall'))>0 else [])),
        ('4. Distribution Details' if not cn else '4. Distribution Details｜分红详情',
         ('以下金额为现金层面的审核数据，不构成自动分红指令。' if cn else 'The amounts below support cash-based review and do not authorize payment.'),
         [('Available for Distribution Review' if not cn else '可供分红审核金额',money(row.get('Available to Distribute'))),('Distribution Paid This Period' if not cn else '本期实际分红',money(row.get('Distribution Paid')))]),
        ('Appendix' if not cn else 'Appendix｜附录',
         ('数据来源：AppFolio Income Statement、Rent Roll 和人工核对的银行现金数据。未上传或未单独提供的字段会明确标注。' if cn else 'Sources: AppFolio Income Statement, Rent Roll, and bank cash verified by management. Fields not supplied are explicitly identified.'),
         [('Property ID' if not cn else '物业编号',str(row.get('Property ID',''))),('Reporting Month' if not cn else '报告月份',str(row.get('Month (YYYY-MM)',''))),('Bank Balance Verification' if not cn else '银行余额核对',money(cash_verified) if cash_verified is not None else ('未核对' if cn else 'Not verified')),('Property Tax Paid' if not cn else '已付地产税',money(row.get('Property Tax Paid'))),('Insurance Paid' if not cn else '已付保险费',money(row.get('Insurance Paid'))),('Other Cash Expense' if not cn else '其他现金支出',money(row.get('Other Cash Expense'))),('Other Cash Adjustments' if not cn else '其他现金调整',money(row.get('Other Cash Adjustments')))]),
    ]
    return sections

def make_docx(row,p,cn):
    d=Document(); s=d.sections[0];s.top_margin=Inches(.85);s.bottom_margin=Inches(.8);s.left_margin=Inches(.95);s.right_margin=Inches(.95)
    normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(10)
    for nm,sz in [('Title',26),('Heading 1',16)]:
        st=d.styles[nm];st.font.name='Calibri';st.font.size=Pt(sz);st.font.color.rgb=RGBColor.from_string(BLUE)
    d.add_paragraph('MOO HOUSING  /  INVESTOR REPORTING','Subtitle')
    d.add_paragraph('\n\n')
    d.add_heading('INVESTOR FINANCIAL REPORT' if not cn else '投资人财务报告',0)
    d.add_paragraph('Financial & Distribution Summary' if not cn else '经营、财务及分红摘要')
    d.add_paragraph('\n')
    d.add_heading(str(p.get('Property Name / LLC','')),level=1)
    d.add_paragraph(str(p.get('Property Address','')))
    d.add_paragraph(('REPORTING PERIOD  /  ' if not cn else '报告期间  /  ')+str(row.get('Month (YYYY-MM)','')))
    d.add_paragraph('\n\nCONFIDENTIAL  •  FOR INVESTOR USE')
    d.add_page_break()
    for idx,(heading,body,items) in enumerate(content(row,p,cn)):
        if idx in (2,4):d.add_page_break()
        d.add_heading(heading,level=1)
        for paragraph in body.split('\n'):
            if paragraph.strip():d.add_paragraph(paragraph)
        if items:
            d.add_paragraph('')
            t=d.add_table(rows=0,cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Light Shading Accent 1'
            for k,v in items:
                a,b=t.add_row().cells;a.text=k;b.text=v
                for cell in (a,b):
                    for par in cell.paragraphs:par.paragraph_format.space_after=Pt(5);par.paragraph_format.space_before=Pt(5)
        d.add_paragraph('\n')
    out=io.BytesIO();d.save(out);return out.getvalue()

def make_pdf(row,p,cn):
    font='Helvetica'
    if cn:pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'));font='STSong-Light'
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='MooTitle2',fontName=font,fontSize=23,leading=31,textColor=colors.HexColor('#'+BLUE),spaceAfter=22))
    styles.add(ParagraphStyle(name='MooHead2',fontName=font,fontSize=15,leading=22,textColor=colors.HexColor('#'+BLUE),spaceBefore=18,spaceAfter=18,keepWithNext=True))
    styles.add(ParagraphStyle(name='MooBody2',fontName=font,fontSize=10,leading=18,spaceAfter=13))
    styles.add(ParagraphStyle(name='MooSmall2',fontName=font,fontSize=9,leading=15,textColor=colors.HexColor('#'+MUTED)))
    esc=lambda x:escape(str(x))
    buf=io.BytesIO();doc=SimpleDocTemplate(buf,pagesize=(595,842),leftMargin=57,rightMargin=57,topMargin=55,bottomMargin=55)
    story=[Paragraph('MOO HOUSING  /  INVESTOR REPORTING',styles['MooSmall2']),Spacer(1,90),Paragraph('投资人财务报告' if cn else 'INVESTOR FINANCIAL REPORT',styles['MooTitle2']),Paragraph('经营、财务及分红摘要' if cn else 'Financial & Distribution Summary',styles['MooBody2']),Spacer(1,34),Paragraph(esc(p.get('Property Name / LLC','')),styles['MooHead2']),Paragraph(esc(p.get('Property Address','')),styles['MooBody2']),Paragraph(esc(row.get('Month (YYYY-MM)','')),styles['MooBody2']),Spacer(1,90),Paragraph('CONFIDENTIAL  /  FOR INVESTOR USE',styles['MooSmall2']),PageBreak()]
    for idx,(heading,body,items) in enumerate(content(row,p,cn)):
        if idx in (2,4):story.append(PageBreak())
        story.append(Paragraph(esc(heading),styles['MooHead2']))
        for para in body.split('\n'):
            if para.strip():story.append(Paragraph(esc(para),styles['MooBody2']))
        if items:
            data=[[Paragraph(esc(k),styles['MooBody2']),Paragraph(esc(v),styles['MooBody2'])] for k,v in items]
            t=Table(data,colWidths=[310,171],hAlign='LEFT')
            t.setStyle(TableStyle([('ROWBACKGROUNDS',(0,0),(-1,-1),[colors.white,colors.HexColor('#EFF4F7')]),('TOPPADDING',(0,0),(-1,-1),12),('BOTTOMPADDING',(0,0),(-1,-1),12),('VALIGN',(0,0),(-1,-1),'TOP')]))
            story.append(t)
        story.append(Spacer(1,30))
    doc.build(story);return buf.getvalue()

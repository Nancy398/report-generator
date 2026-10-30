"""Investor reports built from the actual 1438 Word master, not a visual recreation."""
import io
import os
import shutil
import subprocess
import tempfile
from copy import deepcopy
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table

MASTER = Path(__file__).with_name('1438_template.docx')


def _get(row, *names, default='Not provided'):
    for name in names:
        val = row.get(name)
        if val is not None and str(val).strip() != '':
            return val
    return default


def _cash(v):
    try: return '${:,.2f}'.format(float(str(v).replace(',','').replace('$','')))
    except (TypeError, ValueError): return 'Not provided'


def _pct(v):
    try:
        n=float(v)
        return f'{n*100 if n<=1 else n:.1f}%'
    except (TypeError, ValueError): return 'Not provided'


def _replace_para(p, new_text):
    # Retain paragraph formatting and first run's character formatting.
    if p.runs:
        p.runs[0].text = str(new_text)
        for run in p.runs[1:]: run.text = ''
    else: p.add_run(str(new_text))


def _clone_paragraph(doc, source_element, text=None):
    el=deepcopy(source_element)
    doc.element.body.insert(len(doc.element.body)-1, el)
    p=Paragraph(el,doc)
    if text is not None: _replace_para(p,text)
    return p


def _clone_table(doc, source_element, headings, data, total_last=False):
    el=deepcopy(source_element)
    doc.element.body.insert(len(doc.element.body)-1,el)
    table=Table(el,doc)
    while len(table.rows)>1:
        table._tbl.remove(table.rows[-1]._tr)
    template_row=deepcopy(source_element.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')[1])
    for item in data:
        table._tbl.append(deepcopy(template_row))
    for ri, vals in enumerate([headings]+data):
        for ci, val in enumerate(vals):
            if ci>=len(table.rows[ri].cells): continue
            cell=table.rows[ri].cells[ci]
            _replace_para(cell.paragraphs[0],val)
            for extra in cell.paragraphs[1:]: _replace_para(extra,'')
    if total_last and data:
        # Copy the master table's original total-row shading and bold emphasis.
        original_rows=source_element.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
        if len(original_rows)>2:
            last=deepcopy(original_rows[-1])
            table._tbl.replace(table.rows[-1]._tr,last)
            for ci,val in enumerate(data[-1]):
                _replace_para(table.rows[-1].cells[ci].paragraphs[0],val)
    return table


def make_docx(row,p,cn=False):
    if not MASTER.exists(): raise FileNotFoundError('1438_template.docx must be beside reporting.py')
    doc=Document(MASTER)
    originals=[deepcopy(e) for e in doc.element.body]
    # Keep the original cover (including logo, fonts and spacing), replace content only.
    prop=str(_get(p,'Property Name / LLC','Property Name',default='PROPERTY NAME'))
    addr=str(_get(p,'Property Address',default='Los Angeles, California'))
    period=str(_get(row,'Month (YYYY-MM)',default=datetime.now().strftime('%Y-%m')))
    try: report_date=datetime.strptime(period[:7],'%Y-%m').strftime('%B %Y')
    except ValueError: report_date=period
    cover={3:'Project Overview, Operations & Distribution Summary',4:prop.upper(),5:addr,7:report_date}
    for i,txt in cover.items(): _replace_para(Paragraph(doc.element.body[i],doc),txt)
    for e in list(doc.element.body)[12:-1]: doc.element.body.remove(e)

    occupancy=_pct(_get(row,'Occupancy'))
    occupied=_get(row,'Occupied Rooms')
    total=_get(row,'Total Rooms')
    try: vacant=str(int(float(total))-int(float(occupied)))
    except (ValueError,TypeError): vacant='Not provided'
    summary=str(_get(row,'Summary','Executive Summary','Notes',default=''))
    if not summary.strip(): summary='[Enter your Executive Summary in the Streamlit Summary field.]'
    rent=_cash(_get(row,'Rent Collected','Rent Collection','Collected Rent'))
    ar=_cash(_get(row,'Outstanding AR','Accounts Receivable','AR'))
    income=_cash(_get(row,'Income','Total Income'))
    expenses=_cash(_get(row,'Expenses','Operating Expenses'))
    noi=_cash(_get(row,'NOI'))
    cash=_cash(_get(row,'Ending Cash'))
    reserve=_cash(_get(row,'Required Reserve'))
    shortfall=_cash(_get(row,'Reserve Shortfall'))
    available=_cash(_get(row,'Available to Distribute'))
    distribution=_cash(_get(row,'Distribution Paid'))
    sections=[
      ('01  EXECUTIVE SUMMARY',summary,[('KEY FINANCIAL HIGHLIGHTS',['FINANCIAL METRIC','AMOUNT'],[
          ('Occupancy Rate',occupancy),('Total Income',income),('Net Operating Income (NOI)',noi),('Available for Distribution Review',available)])]),
      ('02  OCCUPANCY & RENT COLLECTION','Occupancy and collection results for the reporting month.',[
          ('OCCUPANCY & COLLECTION',['METRIC','CURRENT MONTH'],[
            ('Occupancy Rate',occupancy),('Occupied Rooms',str(occupied)),('Vacant Rooms',vacant),('Total Rooms',str(total)),('Rent Collected',rent),('Outstanding Balance',ar)])]),
      ('03  FINANCIAL SUMMARY','Operating results shown below are based on the supplied financial records.',[
          ('OPERATING RESULTS',['ITEM','AMOUNT'],[('Total Income',income),('Operating Expenses',expenses),('Net Operating Income (NOI)',noi)]),
          ('CASH & RESERVES',['ITEM','AMOUNT'],[('Ending Cash Balance',cash),('Required Reserve',reserve),('Reserve Shortfall',shortfall)])]),
      ('04  DISTRIBUTION DETAILS','Distributions are subject to management approval, cash verification and unpaid obligations.',[
          ('DISTRIBUTION OVERVIEW',['CASH POSITION','USD'],[('Ending Cash Balance',cash),('Less: Required Reserve',reserve),('Available for Distribution Review',available),('Distribution Paid This Period',distribution)])]),
      ('APPENDIX','Supporting data and financial statements may be attached separately.',[
          ('SOURCE & RECONCILIATION',['ITEM','DETAILS'],[('Property',prop),('Reporting Period',period),('Verified Ending Bank Balance',_cash(_get(row,'Verified Ending Bank Balance'))),('Property Tax Paid',_cash(_get(row,'Property Tax Paid'))),('Insurance Paid',_cash(_get(row,'Insurance Paid'))),('Other Cash Adjustments',_cash(_get(row,'Other Cash Adjustments')))])])
    ]
    for idx,(heading,body,blocks) in enumerate(sections):
        if idx: doc.add_page_break()
        # The original page-top banner is a body paragraph with embedded logo; clone exactly.
        banner=_clone_paragraph(doc, originals[12])
        for run in banner.runs:
            if '1438 West 37th Drive' in run.text:
                run.text=run.text.replace('1438 West 37th Drive · Los Angeles, CA',f'{prop} · {addr}')
        title=_clone_paragraph(doc, originals[13])
        if len(title.runs)>=2:
            parts=heading.split('  ',1)
            title.runs[0].text=(parts[0]+'  ') if len(parts)==2 else ''
            title.runs[1].text=parts[1] if len(parts)==2 else heading
            for run in title.runs[2:]: run.text=''
        else: _replace_para(title,heading)
        _clone_paragraph(doc, originals[14],body)
        for j,(subtitle,cols,items) in enumerate(blocks):
            _clone_paragraph(doc, originals[16],subtitle)
            _clone_table(doc,originals[17],cols,items,total_last=(idx in (0,2,3) and j==len(blocks)-1))
            _clone_paragraph(doc, originals[18], '')
    output=io.BytesIO();doc.save(output);return output.getvalue()


def make_pdf(row,p,cn=False):
    # Use Word-to-PDF conversion, preserving the exact template. This requires LibreOffice.
    binary=shutil.which('libreoffice') or shutil.which('soffice')
    if not binary:
        raise RuntimeError('PDF conversion requires LibreOffice. Download the Word report or deploy with LibreOffice support.')
    with tempfile.TemporaryDirectory() as tmp:
        input_path=Path(tmp)/'investor_report.docx'
        input_path.write_bytes(make_docx(row,p,cn))
        proc=subprocess.run([binary,'-env:UserInstallation=file://'+tmp+'/lo_profile','--headless','--convert-to','pdf','--outdir',tmp,str(input_path)],capture_output=True,text=True,timeout=90)
        pdf=Path(tmp)/'investor_report.pdf'
        if proc.returncode or not pdf.exists(): raise RuntimeError(f'Word-to-PDF conversion failed: {proc.stderr}')
        return pdf.read_bytes()
